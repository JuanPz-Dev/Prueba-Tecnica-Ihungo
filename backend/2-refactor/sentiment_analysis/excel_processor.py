import logging

import openpyxl

from .provider import SentimentProvider

logger = logging.getLogger(__name__)

class ExcelProcessor:
    def __init__(self, provider: SentimentProvider, input_file: str, 
                output_file: str, sheet_name: str, text_column: int,) -> None:
        self.provider = provider
        self.input_file = input_file
        self.output_file = output_file
        self.sheet_name = sheet_name
        self.text_column = text_column

    def process(self) -> None:
        workbook = openpyxl.load_workbook(self.input_file)
        if self.sheet_name not in workbook.sheetnames:
            raise ValueError (f"La hoja '{self.sheet_name}' no existe en el archivo.")

        sheet = workbook[self.sheet_name]
        sheet.cell(1, 4).value = 'NEGATIVO'
        sheet.cell(1, 5).value = 'NEUTRAL'
        sheet.cell(1, 6).value = 'POSITIVO'

        for row in range(2, sheet.max_row + 1):
            text = sheet.cell(row, self.text_column).value
            if text is None:
                logger.warning("La fila %s no contiene texto.", row)
                continue

            result = self.provider.analyze(str(text))
            sheet.cell(row, 4).value = round(result.negative, 3)
            sheet.cell(row, 5).value = round(result.neutral, 3)
            sheet.cell(row, 6).value = round(result.positive, 3)

        workbook.save(self.output_file)
        logger.info("Archivo procesado correctamente: %s", self.output_file)