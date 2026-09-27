from openpyxl import Workbook, load_workbook

from sentiment_analysis.excel_processor import ExcelProcessor
from sentiment_analysis.models import SentimentResult


class FakeSentimentProvider:
    def analyze(self, text: str) -> SentimentResult:
        return SentimentResult(
            negative=10.0,
            neutral=20.0,
            positive=70.0,
        )

def test_process_excel(tmp_path):
    input_file = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    sheet.cell(1, 3).value = "TEXTO"
    sheet.cell(2, 3).value = "Este producto me gustó."

    workbook.save(input_file)

    provider = FakeSentimentProvider()

    processor = ExcelProcessor(
        provider=provider,
        input_file=str(input_file),
        output_file=str(output_file),
        sheet_name="Sheet1",
        text_column=3,
    )

    processor.process()

    result_workbook = load_workbook(output_file)
    result_sheet = result_workbook["Sheet1"]

    assert result_sheet.cell(1, 4).value == "NEGATIVO"
    assert result_sheet.cell(1, 5).value == "NEUTRAL"
    assert result_sheet.cell(1, 6).value == "POSITIVO"

    assert result_sheet.cell(2, 4).value == 10.0
    assert result_sheet.cell(2, 5).value == 20.0
    assert result_sheet.cell(2, 6).value == 70.0

def test_process_ignores_empty_text(tmp_path, caplog):
    input_file = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    sheet.cell(1, 3).value = "TEXTO"
    sheet.cell(2, 3).value = None

    workbook.save(input_file)

    provider = FakeSentimentProvider()
    processor = ExcelProcessor(
        provider=provider,
        input_file=str(input_file),
        output_file=str(output_file),
        sheet_name="Sheet1",
        text_column=3,
    )
    with caplog.at_level("WARNING"):
        processor.process()
    
def test_process_raises_error_for_invalid_sheet(tmp_path):
    input_file = tmp_path / "input.xlsx"
    output_file = tmp_path / "output.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    workbook.save(input_file)

    provider = FakeSentimentProvider()

    processor = ExcelProcessor(
        provider=provider,
        input_file=str(input_file),
        output_file=str(output_file),
        sheet_name="NoExiste",
        text_column=3,
    )

    try:
        processor.process()
    except ValueError as error:
        assert "no existe" in str(error)
    else:
        raise AssertionError("Se esperaba un ValueError")